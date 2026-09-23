"""Fail-fast publication audit for the complete analytics portfolio."""

from __future__ import annotations

import json
import re
import subprocess
import sys
import warnings
import xml.etree.ElementTree as ET
from pathlib import Path
from urllib.parse import unquote

import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl.worksheet._reader")
EXPECTED_PROJECTS = {
    "ecommerce_customer_segmentation_rfm",
    "marketing_campaign_funnel_analysis",
    "ecommerce_growth_conversion_power_bi",
    "commercial_crm_excel_analytics",
    "insurance_risk_segmentation_sql",
    "product_affinity_network_analysis",
    "airline_performance_visual_analytics",
    "supervised_learning_regression_classification_r",
}


def run(command: list[str], cwd: Path) -> None:
    completed = subprocess.run(command, cwd=cwd, text=True, capture_output=True)
    if completed.returncode != 0:
        raise AssertionError(
            f"Command failed: {' '.join(command)}\n{completed.stdout}\n{completed.stderr}"
        )
    print(completed.stdout.strip())


def check_structure() -> None:
    projects = {
        path.name
        for path in ROOT.iterdir()
        if path.is_dir() and path.name not in {"scripts", ".git"}
    }
    assert projects == EXPECTED_PROJECTS, f"Unexpected project set: {sorted(projects)}"

    forbidden_parts = {".pbi", "__pycache__", ".ipynb_checkpoints"}
    forbidden = []
    oversized = []

    # In a cloned Git repository, audit what can actually be published/tracked.
    if (ROOT / ".git").exists():
        completed = subprocess.run(
            ["git", "ls-files", "-z"],
            cwd=ROOT,
            text=False,
            capture_output=True,
        )
        assert completed.returncode == 0, completed.stderr.decode(errors="replace")
        tracked = [
            ROOT / raw.decode("utf-8", errors="replace")
            for raw in completed.stdout.split(b"\0")
            if raw
        ]
        for path in tracked:
            rel_parts = path.relative_to(ROOT).parts
            if any(part in forbidden_parts for part in rel_parts):
                forbidden.append(path)
            if any(part.lower().endswith("_final") for part in rel_parts):
                forbidden.append(path)
            if path.is_file() and path.stat().st_size >= 100 * 1024 * 1024:
                oversized.append(path)
    else:
        # For a publication ZIP/snapshot without .git, scan the filesystem.
        for path in ROOT.rglob("*"):
            rel_parts = path.relative_to(ROOT).parts
            if any(part in forbidden_parts for part in rel_parts):
                forbidden.append(path)
            if path.is_dir() and path.name.lower().endswith("_final"):
                forbidden.append(path)
            if path.is_file() and path.stat().st_size >= 100 * 1024 * 1024:
                oversized.append(path)

    assert not forbidden, f"Forbidden tracked/cache/duplicate paths: {forbidden}"
    assert not oversized, f"Files exceed GitHub's 100 MiB limit: {oversized}"


def check_markdown_links_and_text() -> None:
    link_pattern = re.compile(r"!?\[[^\]]*\]\(([^)]+)\)")
    broken = []
    suspicious = []
    for markdown in ROOT.rglob("*.md"):
        text = markdown.read_text(encoding="utf-8", errors="replace")
        if re.search(r"(?:[A-Za-z]:\\Users\\|/Users/|/home/[^/]+/)", text):
            suspicious.append(markdown)
        for raw_target in link_pattern.findall(text):
            target = raw_target.strip().split("#", 1)[0].strip()
            if not target or target.startswith(("http://", "https://", "mailto:")):
                continue
            candidate = (markdown.parent / unquote(target)).resolve()
            try:
                candidate.relative_to(ROOT.resolve())
            except ValueError:
                broken.append((markdown, raw_target, "escapes repository"))
                continue
            if not candidate.exists():
                broken.append((markdown, raw_target, "missing"))
    assert not broken, f"Broken local Markdown links: {broken}"
    assert not suspicious, f"Machine-specific paths detected: {suspicious}"


def check_python_and_notebooks() -> None:
    for script in ROOT.rglob("*.py"):
        compile(script.read_text(encoding="utf-8"), str(script), "exec")
    for notebook in ROOT.rglob("*.ipynb"):
        content = json.loads(notebook.read_text(encoding="utf-8"))
        code_cells = [cell for cell in content.get("cells", []) if cell.get("cell_type") == "code"]
        errors = [
            output
            for cell in content.get("cells", [])
            for output in cell.get("outputs", [])
            if output.get("output_type") == "error"
        ]
        assert not errors, f"Notebook contains error output: {notebook}"
        assert code_cells and all(cell.get("execution_count") is not None for cell in code_cells), (
            f"Notebook contains unexecuted code cells: {notebook}"
        )


def check_native_and_supporting_artifacts() -> None:
    workbook_path = ROOT / "commercial_crm_excel_analytics/Commercial_CRM_Pipeline_Analytics.xlsx"
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        workbook_formula = load_workbook(workbook_path, read_only=True, data_only=False)
        workbook_values = load_workbook(workbook_path, read_only=True, data_only=True)
    assert workbook_formula.sheetnames == [
        "Dashboard", "CRM_Data", "Rep_Performance", "Channel_Performance",
        "Funnel_Analysis", "Monthly_Trend", "Action_List", "Lookups",
        "Data_Dictionary", "Project_Overview",
    ]
    formula_count = sum(
        1
        for sheet in workbook_formula.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.data_type == "f"
    )
    assert formula_count > 25_000
    excel_errors = {"#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!"}
    displayed_errors = [
        (sheet.title, cell.coordinate, cell.value)
        for sheet in workbook_values.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value in excel_errors
    ]
    assert not displayed_errors, f"Excel displayed errors: {displayed_errors[:10]}"

    powerbi_root = ROOT / "ecommerce_growth_conversion_power_bi/powerbi"
    assert (powerbi_root / "Ecommerce_Growth_Conversion_Analytics.pbip").is_file()
    for json_file in powerbi_root.rglob("*.json"):
        json.loads(json_file.read_text(encoding="utf-8"))

    tableau = ROOT / "airline_performance_visual_analytics/tableau/airline_visual_analytics.twb"
    ET.parse(tableau)
    airline = pd.read_csv(
        ROOT / "airline_performance_visual_analytics/data/processed/airline_performance_analysis.csv",
        low_memory=False,
    )
    assert airline.shape == (80_972, 19)

    sql_text = (
        ROOT / "insurance_risk_segmentation_sql/insurance_charges_analysis.sql"
    ).read_text(encoding="utf-8").upper()
    for token in ("CREATE OR REPLACE VIEW", "WITH ", "PERCENTILE_CONT", "DENSE_RANK", "NULLIF"):
        assert token in sql_text, f"SQL evidence missing: {token}"


def check_flagship_outputs() -> None:
    lifecycle = pd.read_csv(
        ROOT / "ecommerce_customer_segmentation_rfm/outputs/lifecycle_summary.csv"
    )
    retention = pd.read_csv(
        ROOT / "ecommerce_customer_segmentation_rfm/outputs/retention_kpis.csv"
    )
    assert lifecycle["Customers"].sum() == 5_878
    assert abs(lifecycle["Revenue Share (%)"].sum() - 100) < 0.01
    assert set(retention["Month"]) == {1, 3, 6, 12}
    assert retention["Weighted Retention Rate"].between(0, 1).all()

    campaign_quality = pd.read_csv(
        ROOT / "marketing_campaign_funnel_analysis/outputs/data_quality_summary.csv"
    )
    campaign_channels = pd.read_csv(
        ROOT / "marketing_campaign_funnel_analysis/outputs/channel_performance.csv"
    )
    assert campaign_quality["Count"].sum() == 0
    assert campaign_channels["ROAS"].gt(0).all()
    campaign_overall = pd.read_csv(
        ROOT / "marketing_campaign_funnel_analysis/outputs/overall_metrics.csv"
    ).set_index("Metric")["Value"]
    assert int(float(campaign_overall["Campaigns"])) == 10_000
    assert abs(float(campaign_overall["CTR"]) - 0.0547678863) < 1e-8
    assert abs(float(campaign_overall["Click to lead"]) - 0.3012909964) < 1e-8
    assert abs(float(campaign_overall["Lead to conversion"]) - 0.4022744847) < 1e-8

    experiment_summary = pd.read_csv(
        ROOT / "ecommerce_growth_conversion_power_bi/experiment/outputs/experiment_summary.csv"
    ).set_index("Metric")["Value"]
    balance = pd.read_csv(
        ROOT / "ecommerce_growth_conversion_power_bi/experiment/outputs/balance_checks.csv"
    )
    guardrails = pd.read_csv(
        ROOT / "ecommerce_growth_conversion_power_bi/experiment/outputs/guardrail_results.csv"
    )
    assert int(experiment_summary["Users"]) == 32_000
    assert experiment_summary["SRM p-value"] > 0.05
    assert experiment_summary["95% CI low"] > 0
    assert experiment_summary["Primary p-value"] < 0.05
    assert balance["Pass"].all()
    assert guardrails["Pass"].all()


def check_recruiter_contract() -> None:
    root_readme = (ROOT / "README.md").read_text(encoding="utf-8")
    required_phrases = [
        "Customer & Growth Analytics",
        "Start here: three flagship decisions",
        "Customer Lifecycle, Retention & CRM Prioritisation",
        "Campaign & Growth Performance",
        "Product Funnel, Growth & Experimentation",
        "Supporting evidence",
        "Supervised Learning — Regression & Classification",
    ]
    for phrase in required_phrases:
        assert phrase in root_readme, f"Root README missing recruiter cue: {phrase}"


def main() -> None:
    check_structure()
    check_markdown_links_and_text()
    check_python_and_notebooks()
    check_native_and_supporting_artifacts()

    run(
        [sys.executable, "analysis/run_analysis.py"],
        ROOT / "ecommerce_customer_segmentation_rfm",
    )
    run(
        [sys.executable, "analysis/run_analysis.py"],
        ROOT / "marketing_campaign_funnel_analysis",
    )
    run(
        [sys.executable, "experiment/run_experiment.py"],
        ROOT / "ecommerce_growth_conversion_power_bi",
    )
    run(
        [sys.executable, "scripts/validate_dataset.py"],
        ROOT / "ecommerce_growth_conversion_power_bi",
    )

    check_flagship_outputs()
    check_recruiter_contract()
    print("PASS — portfolio publication audit completed")


if __name__ == "__main__":
    main()
